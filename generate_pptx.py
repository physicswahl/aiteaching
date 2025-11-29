"""
Generate PowerPoint presentation from Django demo pages.

This script captures screenshots of each page and creates a PowerPoint presentation.
Requires: pip install selenium pillow python-pptx webdriver-manager

Usage: python generate_pptx.py
"""

from selenium import webdriver
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from webdriver_manager.chrome import ChromeDriverManager
from pptx import Presentation
from pptx.util import Inches
from PIL import Image
import time
import io
import os

# Configuration
BASE_URL = "http://127.0.0.1:8000/demo"

# Pages to capture with optional video timestamps (in navigation order)
# Navigation flow: start -> student_plot -> two_nodes -> video_explanation -> 
# conclusions_one -> five_nodes -> species (external) -> ocr_image -> transition -> 
# word_vectors -> anthropic -> llm -> conclusions_two -> teaching
PAGES = [
    {"path": "/start/", "name": "Start"},
    {"path": "/student-plot/", "name": "Student Plot"},
    {"path": "/two-nodes/", "name": "Two Nodes"},
    {"path": "/video-explanation/", "name": "Video Explanation", "video_time": "0:00.00"},
    {"path": "/conclusions-one/", "name": "Conclusions One"},
    {"path": "/five-nodes/", "name": "Five Nodes"},
    # species page is external to demo app, skipping
    {"path": "/ocr-image/", "name": "OCR Image"},
    {"path": "/transition/", "name": "Transition"},
    {"path": "/word-vectors/", "name": "Word Vectors", "video_time": "0:00.00"},
    # {"path": "/anthropic/", "name": "Anthropic"},  # Excluded
    {"path": "/llm/", "name": "LLM", "video_time": "0:34.50"},
    {"path": "/conclusions-two/", "name": "Conclusions Two"},
    {"path": "/teaching/", "name": "Teaching"},
]

# Pages to exclude
EXCLUDE_PAGES = ["/anthropic/"]


def parse_video_time(time_str):
    """Convert time string (M:SS.CC) to seconds."""
    parts = time_str.split(':')
    minutes = int(parts[0])
    sec_parts = parts[1].split('.')
    seconds = int(sec_parts[0])
    centiseconds = int(sec_parts[1]) if len(sec_parts) > 1 else 0
    return minutes * 60 + seconds + centiseconds / 100


def setup_driver():
    """Setup Chrome WebDriver with appropriate options."""
    chrome_options = Options()
    chrome_options.add_argument('--headless=new')  # Run in headless mode
    chrome_options.add_argument('--window-size=1400,900')
    chrome_options.add_argument('--disable-gpu')
    chrome_options.add_argument('--no-sandbox')
    chrome_options.add_argument('--disable-dev-shm-usage')
    
    service = Service(ChromeDriverManager().install())
    driver = webdriver.Chrome(service=service, options=chrome_options)
    return driver


def capture_page(driver, url, page_name, video_time=None):
    """Capture a screenshot of the page."""
    print(f"Capturing: {page_name}...")
    
    driver.get(url)
    time.sleep(2)  # Wait for page to load
    
    # If there's a video, set the time
    if video_time:
        try:
            # Wait for video element to be present
            video = WebDriverWait(driver, 10).until(
                EC.presence_of_element_located((By.TAG_NAME, "video"))
            )
            
            # Wait a bit more for video to be ready
            time.sleep(1)
            
            time_in_seconds = parse_video_time(video_time)
            
            # Set video time and pause using JavaScript
            driver.execute_script(f"""
                var video = arguments[0];
                video.pause();
                video.currentTime = {time_in_seconds};
                
                // Wait for seek to complete
                return new Promise((resolve) => {{
                    video.onseeked = () => resolve();
                    if (video.readyState >= 2) {{
                        resolve();
                    }}
                }});
            """, video)
            
            # Extra wait to ensure frame is rendered
            time.sleep(2)
            print(f"  Video set to {video_time}")
        except Exception as e:
            print(f"  Warning: Could not set video time: {e}")
    
    # Additional wait for any animations/rendering
    time.sleep(1)
    
    # Take screenshot
    screenshot = driver.get_screenshot_as_png()
    return Image.open(io.BytesIO(screenshot))


def create_presentation(screenshots):
    """Create PowerPoint presentation from screenshots."""
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(5.625)  # 16:9 aspect ratio
    
    for name, img in screenshots:
        print(f"Adding slide: {name}")
        slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout
        
        # Save image temporarily
        temp_img_path = f"temp_slide_{name.replace(' ', '_')}.png"
        img.save(temp_img_path)
        
        # Add image to slide
        slide.shapes.add_picture(temp_img_path, 0, 0, width=prs.slide_width, height=prs.slide_height)
        
        # Clean up temp file
        os.remove(temp_img_path)
    
    return prs


def main():
    print("Starting PowerPoint generation...")
    print("=" * 50)
    
    # Filter out excluded pages
    pages_to_capture = [p for p in PAGES if p["path"] not in EXCLUDE_PAGES]
    
    driver = None
    try:
        driver = setup_driver()
        screenshots = []
        
        for page in pages_to_capture:
            url = BASE_URL + page["path"]
            video_time = page.get("video_time")
            
            img = capture_page(driver, url, page["name"], video_time)
            screenshots.append((page["name"], img))
        
        print("\n" + "=" * 50)
        print("Creating PowerPoint presentation...")
        prs = create_presentation(screenshots)
        
        output_file = "AI_Teaching_Presentation.pptx"
        prs.save(output_file)
        
        print(f"\n✓ PowerPoint saved as: {output_file}")
        print(f"  Total slides: {len(screenshots)}")
        
    except Exception as e:
        print(f"\n✗ Error: {e}")
        import traceback
        traceback.print_exc()
    
    finally:
        if driver:
            driver.quit()


if __name__ == "__main__":
    main()
